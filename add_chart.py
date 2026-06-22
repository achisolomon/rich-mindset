# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, RadarChart
from openpyxl.chart.label import DataLabelList

p = "/sessions/exciting-peaceful-bardeen/mnt/outputs/שאלון_מיפוי_מיינדסט.xlsx"
wb = load_workbook(p)
ws2 = wb["תוצאות"]

# categories rows 3..15, total row 16
first, last = 3, 15
data = Reference(ws2, min_col=2, min_row=2, max_row=last)   # include header row 2 for title
cats = Reference(ws2, min_col=1, min_row=first, max_row=last)

# Horizontal bar chart
bar = BarChart()
bar.type = "bar"
bar.title = "פרופיל מיינדסט לפי קטגוריה (1=עני · 5=עשיר)"
bar.add_data(data, titles_from_data=True)
bar.set_categories(cats)
bar.y_axis.scaling.min = 0
bar.y_axis.scaling.max = 5
bar.x_axis.delete = False
bar.y_axis.delete = False
bar.height = 12
bar.width = 24
bar.dLbls = DataLabelList(); bar.dLbls.showVal = True
bar.legend = None
ws2.add_chart(bar, "E2")

# Radar chart on its own area
radar = RadarChart()
radar.type = "filled"
radar.title = "מכ\"ם מיינדסט"
radar.add_data(data, titles_from_data=True)
radar.set_categories(cats)
radar.y_axis.scaling.min = 0
radar.y_axis.scaling.max = 5
radar.height = 12
radar.width = 14
radar.legend = None
ws2.add_chart(radar, "E26")

wb.save(p)
print("ok")
